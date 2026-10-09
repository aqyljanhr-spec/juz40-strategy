"""JUZ40 Strategy — FastAPI/SQLite standalone event session server."""
import os,json,sqlite3,threading,datetime,io,zipfile,urllib.request
import psycopg
from pathlib import Path
from fastapi import FastAPI,HTTPException,WebSocket,WebSocketDisconnect,Request
from fastapi.responses import FileResponse,StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
from openpyxl import Workbook
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

ROOT=Path(__file__).parent
STAFF=json.loads((ROOT/'staff.json').read_text())
GROUPS=list(dict.fromkeys(p['group'] for p in STAFF))
DEFAULT_CAPTAINS={g:next(x['id'] for x in STAFF if x['group']==g) for g in GROUPS}
DB=Path(os.environ.get('JUZ40_DB',str(ROOT/'strategy.sqlite3')))
DATABASE_URL=os.environ.get('DATABASE_URL','')
app=FastAPI(title='JUZ40 Strategy Platform')
lock=threading.RLock()
clients={'test':set(),'live':set()}

class Action(BaseModel):
    mode:str='test'
    actor:int=0
    action:str
    data:dict={}

def now(): return datetime.datetime.now(datetime.timezone.utc).isoformat()
def default():
 return {'staff':[{**p,'captain':p['id']==DEFAULT_CAPTAINS[p['group']]} for p in STAFF], 'joined':[], 'opened':[False]*4,'closed':[False]*4,'phase':'writing','issues':[],'votes':{},'shared':{},'finished':{},'session_closed':False,'reports':[],'audit':[],'revision':0}
def connection():
 if DATABASE_URL:
  c=psycopg.connect(DATABASE_URL)
  c.execute('CREATE TABLE IF NOT EXISTS state (mode TEXT PRIMARY KEY,payload TEXT NOT NULL)')
  c.commit()
  return c
 DB.parent.mkdir(parents=True,exist_ok=True)
 c=sqlite3.connect(str(DB));c.execute('CREATE TABLE IF NOT EXISTS state (mode TEXT PRIMARY KEY,payload TEXT NOT NULL)');return c

def load(mode):
 with connection() as c:
  if DATABASE_URL:r=c.execute('SELECT payload FROM state WHERE mode=%s',(mode,)).fetchone()
  else:r=c.execute('SELECT payload FROM state WHERE mode=?',(mode,)).fetchone()
 return json.loads(r[0]) if r else default()

def save(mode,s):
 with connection() as c:
  if DATABASE_URL:c.execute('INSERT INTO state(mode,payload) VALUES(%s,%s) ON CONFLICT(mode) DO UPDATE SET payload=excluded.payload',(mode,json.dumps(s,ensure_ascii=False)))
  else:c.execute('INSERT INTO state(mode,payload) VALUES(?,?) ON CONFLICT(mode) DO UPDATE SET payload=excluded.payload',(mode,json.dumps(s,ensure_ascii=False)))

def role(s,actor):
 p=next((p for p in s['staff'] if p['id']==actor),None)
 if not p: raise HTTPException(403,'Қызметкер тізімде жоқ')
 return p

def check_action(s,p,action):
 if action.startswith('admin_'): return # No identity proof: admin access is unprotected by chosen name-only access model
 if s['session_closed'] and action not in ['report_save']:raise HTTPException(409,'Сессия жабық')
 if action not in ['join','report_save'] and (not s['opened'][{'issue_add':0,'issue_edit':0,'vote':0,'shared':int(action.split(':')[1]) if ':' in action else 0,'finish':int(action.split(':')[1]) if ':' in action else 0}.get(action,0)]):raise HTTPException(409,'Блок жабық')

def mutate(s,p,action,d):
 uid=p['id'];group=p['group'];phase=s['phase']
 if action=='join':
  if uid not in s['joined']:s['joined'].append(uid)
 elif action=='issue_add':
  if phase!='writing':raise HTTPException(409,'Мәселе жазу кезеңі аяқталды')
  txt=str(d.get('text','')).strip()
  if not 3<=len(txt)<=3000:raise HTTPException(400,'Мәселені 3–3000 таңба аралығында жазыңыз')
  if len([x for x in s['issues'] if x['author']==uid])>=3:raise HTTPException(400,'Ең көбі 3 мәселе')
  n=max([x['id'] for x in s['issues']],default=0)+1;s['issues'].append({'id':n,'author':uid,'group':group,'text':txt})
 elif action=='issue_edit':
  if phase!='writing':raise HTTPException(409,'Мәселелер бекітілген')
  issue=next((x for x in s['issues'] if x['id']==int(d.get('id',-1)) and x['author']==uid),None)
  if not issue:raise HTTPException(404,'Мәселе табылмады')
  txt=str(d.get('text','')).strip()
  if len(txt)<3 or len(txt)>3000:raise HTTPException(400,'3–3000 таңба қажет')
  issue['text']=txt
 elif action=='vote':
  if phase!='voting':raise HTTPException(409,'Дауыс беру кезеңі ашық емес')
  ids=d.get('ids',[])
  if len(ids)!=3 or len(set(ids))!=3 or any(not any(x['id']==i and x['group']==group for x in s['issues']) for i in ids):raise HTTPException(400,'Өз тобыңыздың дәл 3 түрлі мәселесін таңдаңыз')
  s['votes'][str(uid)]=ids
 elif action.startswith('shared:'):
  block=int(action.split(':')[1]);key=str(d.get('key',''))
  if not p.get('captain'):raise HTTPException(403,'Тек капитан жаза алады')
  if s['closed'][block] or s['finished'].get(group+'|'+str(block)):raise HTTPException(409,'Блок өңдеуге жабық')
  allowed={0:['solution_0','solution_1','solution_2','choice_2'],1:['department_1','department_3','department_5','company_1','company_3','company_5','publisher'],2:['answer'],3:['answer']}
  if key not in allowed[block] or key=='publisher' and p['dept']!='Әдістеме':raise HTTPException(400,'Қате өріс')
  if block==0 and phase!='final':raise HTTPException(409,'Қорытынды кезеңі ашылмаған')
  text=str(d.get('text',''))
  if len(text)>20000:raise HTTPException(400,'Мәтін тым ұзын')
  s['shared'][group+'|'+str(block)+'|'+key]=text
 elif action.startswith('finish:'):
  b=int(action.split(':')[1]);
  if not p.get('captain'):raise HTTPException(403,'Тек капитан')
  if s['closed'][b]:raise HTTPException(409,'Блок жабық')
  if b==0 and s['phase']!='final':raise HTTPException(409,'Қорытынды кезеңі ашылмаған')
  s['finished'][group+'|'+str(b)]=True
 elif action=='admin_stage':
  phase2=d.get('phase');
  if phase2 not in ['writing','voting','final']:raise HTTPException(400,'Қате кезең')
  s['phase']=phase2
 elif action=='admin_block':
  b=int(d['block']);s['opened'][b]=bool(d['open']);s['closed'][b]=bool(d.get('close',False)) if not d['open'] else False
 elif action=='admin_reopen':
  b=int(d['block']);g=str(d['group']);s['finished'].pop(g+'|'+str(b),None);s['closed'][b]=False;s['opened'][b]=True
 elif action=='admin_assign':
  target=role(s,int(d['id']));g=str(d['group'])
  if g not in GROUPS or g.split(' · ')[0]!=target['dept']:raise HTTPException(400,'Топ бөлімге сай емес')
  old=target['group'];target['group']=g
  if target['captain']:target['captain']=False
  if not any(x['captain'] and x['group']==old for x in s['staff']):pass
 elif action=='admin_captain':
  target=role(s,int(d['id']));
  for x in s['staff']:
   if x['group']==target['group']:x['captain']=x['id']==target['id']
 elif action=='admin_edit':
  g=d['group'];b=int(d['block']);key=d['key'];val=str(d['text'])
  if g not in GROUPS or len(val)>20000:raise HTTPException(400,'Қате дерек')
  s['shared'][g+'|'+str(b)+'|'+str(key)]=val
 elif action=='admin_finish':
  s['session_closed']=True;s['opened']=[False]*4;s['closed']=[True]*4
 elif action=='admin_resume':s['session_closed']=False
 elif action=='report_save':
  text=str(d.get('text',''))
  if len(text)>40000:raise HTTPException(400,'Есеп тым ұзын')
  s['reports'].append({'version':len(s['reports'])+1,'text':text,'at':now(),'editor':uid})
 else:raise HTTPException(400,'Белгісіз әрекет')

@app.get('/api/state/{mode}')
def get_state(mode:str, actor:int=0, admin:bool=False):
 if mode not in clients:raise HTTPException(400,'mode')
 s=load(mode)
 if admin:return s # Name-only admin access chosen by the user; not secure identity proof
 if actor:
  p=role(s,actor);g=p['group'];s['issues']=[i for i in s['issues'] if i['group']==g]
  s['votes']={k:v for k,v in s['votes'].items() if any(int(k)==x['id'] and x['group']==g for x in s['staff'])} if s['phase']=='final' else {str(actor):s['votes'].get(str(actor),[])}
  s['shared']={k:v for k,v in s['shared'].items() if k.startswith(g+'|')}
  s['reports']=[];s['audit']=[]
 return s

@app.post('/api/action')
async def post_action(payload:Action):
 if payload.mode not in clients:raise HTTPException(400,'mode')
 with lock:
  s=load(payload.mode);p=role(s,payload.actor)
  check_action(s,p,payload.action)
  mutate(s,p,payload.action,payload.data)
  s['revision']+=1
  s['audit'].append({'at':now(),'actor':p['name'],'action':payload.action,'data':{k:(v[:100] if isinstance(v,str) else v) for k,v in payload.data.items()}})
  s['audit']=s['audit'][-1000:]
  save(payload.mode,s)
 dead=[]
 for ws in list(clients[payload.mode]):
  try:await ws.send_json({'revision':s['revision']})
  except:dead.append(ws)
 for w in dead:clients[payload.mode].discard(w)
 return {'ok':True,'revision':s['revision']}

@app.websocket('/api/live/{mode}')
async def live(websocket:WebSocket,mode:str):
 if mode not in clients:await websocket.close(code=1008);return
 await websocket.accept();clients[mode].add(websocket)
 try:
  while True:await websocket.receive_text()
 except WebSocketDisconnect:pass
 finally:clients[mode].discard(websocket)

@app.get('/api/export/{mode}.xlsx')
def excel(mode:str):
 if mode not in clients:raise HTTPException(400,'mode')
 s=load(mode);wb=Workbook();ws=wb.active;ws.title='Қатысушылар';ws.append(['Аты-жөні','Бөлім','Топ','Капитан','Check-in'])
 for p in s['staff']:ws.append([p['name'],p['dept'],p['group'],'Иә' if p['captain'] else 'Жоқ','Иә' if p['id'] in s['joined'] else 'Жоқ'])
 w=wb.create_sheet('Мәселелер');w.append(['Топ','Қызметкер','Мәселе','Дауыс'])
 for x in s['issues']:w.append([x['group'],next((p['name'] for p in s['staff'] if p['id']==x['author']),''),x['text'],sum(x['id'] in votes for votes in s['votes'].values())])
 w=wb.create_sheet('Топ жауаптары');w.append(['Топ','Блок','Өріс','Жауап'])
 for k,v in s['shared'].items():
  g,b,f=k.split('|',2);w.append([g,int(b)+1,f,v])
 w=wb.create_sheet('AI есеп нұсқалары');w.append(['Нұсқа','Уақыты','Мәтін'])
 for r in s['reports']:w.append([r['version'],r['at'],r['text']])
 for sh in wb:
  sh.freeze_panes='A2';sh.auto_filter.ref=sh.dimensions
  for col in sh.columns:
   letter=col[0].column_letter;sh.column_dimensions[letter].width=min(58,max(16,max(len(str(c.value or '')) for c in col[:50])+3))
 data=io.BytesIO();wb.save(data);data.seek(0)
 return StreamingResponse(data,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="juz40-{mode}.xlsx"'})

@app.get('/api/team/{mode}/{group_no}/pdf')
def pdf(mode:str,group_no:int,actor:int):
 if mode not in clients:raise HTTPException(400,'mode')
 s=load(mode);p=role(s,actor);g=p['group'];
 if GROUPS.index(g)!=group_no:raise HTTPException(403,'Тек өз тобыңыз')
 font='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
 if os.path.isfile(font):pdfmetrics.registerFont(TTFont('DejaVu',font))
 buf=io.BytesIO();c=canvas.Canvas(buf,pagesize=A4);c.setFont('DejaVu' if os.path.isfile(font) else 'Helvetica',11)
 y=790
 def line(t):
  nonlocal y
  for chunk in str(t).split('\n'):
   while chunk:
    part=chunk[:75];chunk=chunk[75:]
    if y<65:c.showPage();c.setFont('DejaVu' if os.path.isfile(font) else 'Helvetica',11);y=790
    c.drawString(35,y,part);y-=17
 c.setTitle('JUZ40 Strategy · '+g)
 line('JUZ40 STRATEGY · '+g)
 titles=['1. Қазіргі жағдай','2. Vision','3. Тиімділік','4. ҰБТ және тәуекелдер']
 for b,title in enumerate(titles):
  y-=9;line(title)
  for k,v in s['shared'].items():
   if k.startswith(g+'|'+str(b)+'|'):line(k.split('|',2)[2]+': '+v)
 c.save();buf.seek(0)
 return StreamingResponse(buf,media_type='application/pdf',headers={'Content-Disposition':f'attachment; filename="juz40-group-{group_no+1}.pdf"'})

class AIRequest(BaseModel):
 mode:str='live'
 scope:str='all'
 kind:str='strategic'
 actor:int=0
@app.post('/api/ai')
def ai(req:AIRequest):
 if not os.environ.get('OPENAI_API_KEY'):raise HTTPException(503,'AI API кілті орнатылмаған. OPENAI_API_KEY қажет.')
 s=load(req.mode);role(s,req.actor)
 group_names=GROUPS if req.scope=='all' else [g for g in GROUPS if g.startswith(req.scope+' ·')]
 answers={k:v for k,v in s['shared'].items() if any(k.startswith(g+'|') for g in group_names)}
 prompt='JUZ40 стратегиялық сессиясы. Қазақ тілінде '+req.kind+' талдау жасаңыз. Тек төмендегі бастапқы жауаптарға сүйеніңіз; қолдау таппаған тұжырымдарды гипотеза деп белгілеңіз; әр ұсыныстың дереккөз-тобын көрсетіңіз; нақты адам мен KPI ойдан қоспаңыз.\n'+json.dumps(answers,ensure_ascii=False)
 body={'model':os.getenv('OPENAI_MODEL','gpt-4.1-mini'),'messages':[{'role':'user','content':prompt}],'temperature':0.2}
 request=urllib.request.Request('https://api.openai.com/v1/chat/completions',data=json.dumps(body).encode(),headers={'Authorization':'Bearer '+os.environ['OPENAI_API_KEY'],'Content-Type':'application/json'})
 try:
  with urllib.request.urlopen(request,timeout=60) as res:data=json.load(res)
  return {'report':data['choices'][0]['message']['content']}
 except Exception as e:raise HTTPException(502,'AI қызметі жауап бермеді: '+str(e)[:180])

app.mount('/',StaticFiles(directory=str(ROOT/'static'),html=True),name='static')
